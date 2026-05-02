import json, os
from datetime import datetime

emails_data = [
    {"cnj": "0004037-40.2015.8.08.0004", "tribunal": "TJES", "qtd": 1, "ultima_data": "2026-05-01"},
    {"cnj": "0003454-21.2016.8.08.0004", "tribunal": "TJES", "qtd": 3, "ultima_data": "2026-05-01"},
    {"cnj": "5004093-24.2021.8.08.0021", "tribunal": "TJES", "qtd": 3, "ultima_data": "2026-04-30"},
    {"cnj": "0000783-88.2017.8.08.0004", "tribunal": "TJES", "qtd": 5, "ultima_data": "2026-04-30"},
    {"cnj": "0600276-54.2024.6.08.0019", "tribunal": "TRE-ES", "qtd": 5, "ultima_data": "2026-04-29"},
    {"cnj": "5001026-68.2022.8.08.0004", "tribunal": "TJES", "qtd": 1, "ultima_data": "2026-04-29"},
    {"cnj": "5000931-38.2022.8.08.0004", "tribunal": "TJES", "qtd": 2, "ultima_data": "2026-04-29"},
    {"cnj": "0002045-10.2016.8.08.0004", "tribunal": "TJES", "qtd": 1, "ultima_data": "2026-04-27"},
    {"cnj": "0000932-26.2013.8.08.0004", "tribunal": "TJES", "qtd": 1, "ultima_data": "2026-04-27"},
]

def normalize_cnj(cnj):
    return cnj.replace("-","").replace(".","").replace(" ","")

excel_path = r"G:\Meu Drive\Pessoal\Pericia Contabil\Relatorio_Pericias_Atualizado_Atual.xlsx"
planilha = []
try:
    import openpyxl
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = next((wb[s] for s in wb.sheetnames if any(k in s for k in ["Fonte","Laudo","dados"])), wb.active)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0: continue
        try:
            cnj = str(row[2]).strip() if row[2] else ""
            if cnj and cnj != "None":
                planilha.append({"cnj": cnj, "vara": str(row[1] or ""), "autor": str(row[3] or ""),
                                  "status": str(row[8] or ""), "pendente": str(row[11] or "")})
        except: pass
    wb.close()
    print(f"Excel lido: {len(planilha)} registros")
except Exception as e:
    print(f"AVISO: {e}")

urgentes, atencao, informativos, sem_correspondencia = [], [], [], []
for email in emails_data:
    found = next((p for p in planilha if normalize_cnj(p["cnj"]) == normalize_cnj(email["cnj"])), None)
    if not found:
        sem_correspondencia.append(email["cnj"])
        continue
    s = found["status"].lower()
    item = {**email, "autor": found["autor"], "vara": found["vara"],
            "status": found["status"], "pendente": found["pendente"], "qtd_notificacoes": email["qtd"]}
    if "laudo pendente" in s or "quesita" in s or email["qtd"] >= 3: urgentes.append(item)
    elif "laudo entregue" in s or "pendente recebimento" in s: atencao.append(item)
    else: informativos.append(item)

resultado = {
    "data_verificacao": datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
    "total_movimentacoes": sum(e["qtd"] for e in emails_data),
    "total_processos": len(emails_data),
    "excel_lido": len(planilha) > 0,
    "urgentes": urgentes, "atencao": atencao,
    "informativos": informativos, "sem_correspondencia": sem_correspondencia,
}

for path in [r"G:\Meu Drive\Pessoal\Pericia Contabil\agente_pericias\ultima_verificacao_intimacoes.json",
             os.path.join(os.path.dirname(os.path.abspath(__file__)), "ultima_verificacao_intimacoes.json")]:
    try:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(resultado, f, ensure_ascii=False, indent=2)
        print(f"Salvo: {path}")
        break
    except Exception as e:
        print(f"Erro: {e}")

print("=== CONCLUIDO ===")
